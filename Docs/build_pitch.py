"""Create the editable eight-slide pitch. Run from the repository root."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

CREAM = RGBColor(255, 246, 232)
TEAL = RGBColor(14, 125, 116)
INK = RGBColor(31, 42, 46)
WHITE = RGBColor(255, 255, 255)

SLIDES = [
    {
        "kicker": "SE4041 Assignment 02, Part B",
        "title": "Little Discoverers",
        "subtitle": "Look. Choose. Discover.",
        "bullets": [
            "A single-player picture-recognition game for Apple TV.",
            "For children aged 4–7, with a grown-up nearby.",
            "Look at a large picture. Choose its name with the Siri Remote.",
        ],
        "notes": (
            "This is a tvOS prototype, separate from any Part A iOS app. "
            "A child picks Animals, Fruits, or Vehicles, sees one picture, and presses the remote. "
            "It is meant for the sofa, not a phone."
        ),
    },
    {
        "kicker": "Problem and audience",
        "title": "A television is not a touch screen",
        "subtitle": "One remote, a shared screen, and a grown-up nearby.",
        "bullets": [
            "Many early-learning games are built for phones and tablets.",
            "This activity is a short naming session at TV distance.",
            "A wrong try does not take stars away.",
            "No timer and no ranking against other children.",
        ],
        "notes": (
            "Do not claim that television teaches better. "
            "The audience is children aged 4 to 7 and the grown-up with them. "
            "The value in this prototype is a calm five-picture session."
        ),
    },
    {
        "kicker": "Gameplay",
        "title": "Look, choose, then discover",
        "subtitle": "Five pictures. Three answers. One star for a first try.",
        "bullets": [
            "Topics: Animals, Fruits, and Vehicles. Five pictures each.",
            "“What is this?” sits above the picture, beside three large answers.",
            "A wrong choice turns off. The child can try another.",
            "Hear Question reads the choices and does not give the answer.",
            "A finished session saves a best score on this Apple TV.",
        ],
        "notes": (
            "Show one wrong answer and one first-try star. "
            "Leaving early does not save. Finishing saves once. "
            "Mention narration only if you have heard it."
        ),
    },
    {
        "kicker": "Platform",
        "title": "Built for the Siri Remote",
        "subtitle": "Focus, Menu, and a 10-foot layout.",
        "bullets": [
            "Every control is a SwiftUI button with a clear focus border.",
            "After a wrong answer, focus moves to an available choice.",
            "After a correct answer, focus moves to Next Picture.",
            "Menu asks before an unfinished quiz is discarded.",
            "HD and 4K share a 1920 by 1080 point layout.",
        ],
        "notes": (
            "Show the focus move and the leave confirmation. "
            "That is the Apple TV behaviour. It is not a resized phone screen."
        ),
    },
    {
        "kicker": "Similar apps",
        "title": "Video on Apple TV, games on touch screens",
        "subtitle": "Checked against primary App Store and support pages.",
        "bullets": [
            "PBS KIDS Video is on Apple TV. It streams shows. It is not this quiz.",
            "PBS KIDS Games, Khan Academy Kids, and Sago Mini School were listed for iPhone and iPad, not Apple TV.",
            "Little Discoverers combines a remote quiz, offline pictures, spoken choices, and saved progress.",
            "It does not replace those larger libraries.",
        ],
        "notes": (
            "Open the live App Store page before the viva, because the text can change. "
            "Do not claim a tablet app is missing from Apple TV unless you have checked that page."
        ),
    },
    {
        "kicker": "A hypothesis, not a result",
        "title": "Free topics, paid packs later",
        "subtitle": "No purchases are built into this prototype.",
        "bullets": [
            "Possible later model: these three topics stay free.",
            "Extra topic packs could be a one-time purchase.",
            "First step would be a small supervised pilot.",
            "No price has been tested. No revenue is claimed.",
        ],
        "notes": (
            "This is a guess. There is no StoreKit code. "
            "Do not name a price until caregivers have used the free topics."
        ),
    },
    {
        "kicker": "What is actually true",
        "title": "Prototype status",
        "subtitle": "Say what has been built, and what has not been run.",
        "bullets": [
            "Built: six screens, 15 local pictures, scoring, narration code, and on-device progress.",
            "Checked here: the question file, the image files, and that the Xcode project is present.",
            "Not checked here: Xcode build, XCTest, the simulator, speech, VoiceOver, or a real Apple TV.",
        ],
        "notes": (
            "If you have run the app since this slide was written, replace the status with what you saw, including failures. "
            "Do not say the tests passed unless you ran them."
        ),
    },
    {
        "kicker": "Next",
        "title": "Pilot first, store later",
        "subtitle": "Impact is possible. It is not measured.",
        "bullets": [
            "Run the simulator checklist: remote, narration, relaunch, Reduce Motion, and VoiceOver.",
            "Add sharper 4K art, a layered app icon, and Top Shelf before any store conversation.",
            "Ask a few families to try the free topics before designing paid packs.",
            "A fully nonvisual mode is future work, not a claim about this build.",
        ],
        "notes": (
            "Close on the next action. A possible impact is a few calm minutes of shared naming on the family TV. "
            "Do not promise a mark or a learning gain."
        ),
    },
]


def add_text(slide, text, left, top, width, height, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return frame


def build(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for index, item in enumerate(SLIDES, start=1):
        slide = presentation.slides.add_slide(blank)
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), presentation.slide_width, presentation.slide_height
        )
        background.line.fill.background()
        background.fill.solid()
        background.fill.fore_color.rgb = CREAM
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(7.5))
        bar.line.fill.background()
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL
        add_text(slide, f"{index}  /  8", 0.7, 0.28, 2, 0.35, 14, True, TEAL)
        add_text(slide, item["kicker"].upper(), 2.4, 0.28, 8, 0.35, 14, True, TEAL)
        add_text(slide, item["title"], 0.7, 0.75, 12, 0.9, 36, True, INK)
        add_text(slide, item["subtitle"], 0.7, 1.7, 12, 0.7, 22, False, TEAL)
        frame = add_text(slide, item["bullets"][0], 0.7, 2.6, 12, 4.4, 22, False, INK)
        for bullet in item["bullets"][1:]:
            paragraph = frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(22)
            paragraph.font.color.rgb = INK
            paragraph.font.name = "Calibri"
            paragraph.space_after = Pt(12)
        notes = slide.notes_slide.notes_text_frame
        notes.text = item["notes"]
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


if __name__ == "__main__":
    build(Path(__file__).resolve().parent / "LittleDiscoverers-Pitch.pptx")
